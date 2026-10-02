"""Gate: replicate each workbook's Derivation for the US row in Python and compare to its CSV.

Reads only the workbook source tabs and the Regions US row (row 5), never cached formula
values (openpyxl saves drop them). Then compares every InputData/indst CSV's git blob hash
to the snapshot taken before loading (csv_hashes_before.txt in this folder).
Run:  <python> Cement_Regional_Rows_Loading/check_us_rows.py
"""
import csv
import subprocess
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
IND = ROOT / "InputData" / "indst"
fails = []


def wb(rel):
    return openpyxl.load_workbook(IND / rel)


def csv_rows(rel):
    with open(IND / rel, newline="") as f:
        r = list(csv.reader(f))
    return r[0][1:], [float(x) for x in r[1][1:]]


def check(name, expected, actual, tol):
    bad = [(i, e, a) for i, (e, a) in enumerate(zip(expected, actual)) if abs(e - a) > tol]
    ok = len(expected) == len(actual) and not bad
    print(f"{name:10s} {'PASS' if ok else 'FAIL'}  n={len(actual)}" + ("" if ok else f"  first mismatches {bad[:3]}"))
    if not ok:
        fails.append(name)


def us(ws):
    assert ws["A5"].value == "United States"
    return ws


def usgs(w):
    u = w["USGS MCS 2026 Cement"]
    return {k: [u.cell(r, c).value for c in range(2, 7)] for k, r in
            dict(cement=4, clinker=5, clk_imp=8, cons=10).items()}


YEARS = list(range(2021, 2051))
TAU = [round(0.05 * i, 2) for i in range(21)]

# BCD: 2021-25 USGS apparent consumption; 2026+ anchor x AEO index (shipments(y)/shipments(2025))
w = wb("BCD/BAU Cement Demand.xlsx"); R = us(w["Regions"]); s = usgs(w)
aeo = [w["AEO 2026 Table 23"].cell(4, c).value for c in range(5, 31)]  # 2025..2050
exp = [1e6 * s["cons"][y - 2021] / 1000 if y <= 2025 else 1e6 * R["C5"].value * aeo[y - 2025] / aeo[0] for y in YEARS]
check("BCD", exp, csv_rows("BCD/BCD.csv")[1], 1.0)

# BCtCR: 2021-25 round((clinker + clinker imports)/cement, 4); then linear anchor -> target, round 5
w = wb("BCtCR/BAU Clinker to Cement Ratio.xlsx"); R = us(w["Regions"]); s = usgs(w)
hist = [round((s["clinker"][i] + s["clk_imp"][i]) / s["cement"][i], 4) for i in range(5)]
a_y, a, t, t_y = R["B5"].value, R["C5"].value, R["E5"].value, R["F5"].value
exp = [hist[y - 2021] if y <= 2025 else round(a + (t - a) * max(0, min(1, (y - a_y) / (t_y - a_y))), 5) for y in YEARS]
check("BCtCR", exp, csv_rows("BCtCR/BCtCR.csv")[1], 1e-9)

# BFoCDMbNI: round(1 - clinker / (consumption x blend ratio), 4); held at 2025 after
w = wb("BFoCDMbNI/BAU Fraction of Clinker Demand Met by Net Imports.xlsx"); R = us(w["Regions"]); s = usgs(w)
assert R["C5"].value == "=Derivation!G10", "US row of BFoCDMbNI is no longer the Derivation link"
hist = [round(1 - s["clinker"][i] / (s["cons"][i] * round((s["clinker"][i] + s["clk_imp"][i]) / s["cement"][i], 4)), 4) for i in range(5)]
exp = [hist[min(y, 2025) - 2021] for y in YEARS]
check("BFoCDMbNI", exp, csv_rows("BFoCDMbNI/BFoCDMbNI.csv")[1], 1e-9)

# Scalars: Regions US value == CSV value
for name, rel, c in [("MACtCR", "MACtCR/Minimum Achievable Clinker to Cement Ratio.xlsx", "B5"),
                     ("CULB", "CULB/Cement Utilization Lower Bound.xlsx", "B5"),
                     ("CUUB", "CUUB/Cement Utilization Upper Bound.xlsx", "B5")]:
    R = us(wb(rel)["Regions"])
    check(name, [R[c].value], csv_rows(f"{name}/{name}.csv")[1], 1e-12)

# CRBW: w = tau^p / (tau^p + k (1-tau)^p), round 4, endpoints pinned
R = us(wb("CRBW/Cement Replacement Blend Weight vs Tightness.xlsx")["Regions"]); p, k = R["B5"].value, R["C5"].value
exp = [0 if x <= 0 else 1 if x >= 1 else round(x ** p / (x ** p + k * (1 - x) ** p), 4) for x in TAU]
# Tolerance = the workbook's own Checks tab (max abs diff < 0.00015): the CSV was generated with the
# unrounded k (0.368563-0.368585 reproduces every literal); Regions shows k rounded to 0.3686, which
# moves two points (tau 0.40, 0.65) by 0.0001.
check("CRBW", exp, csv_rows("CRBW/CRBWvT.csv")[1], 1.5e-4)

# CTOA: 0 for tau <= tau0, else sat x ((tau - tau0)/(1 - tau0))^2, round 4
R = us(wb("CTOA/Cement Tightness Operating Adder vs Tightness.xlsx")["Regions"]); t0, sat = R["B5"].value, R["C5"].value
exp = [0 if x <= t0 else round(sat * ((x - t0) / (1 - t0)) ** 2, 4) for x in TAU]
check("CTOA", exp, csv_rows("CTOA/CTOAvT.csv")[1], 1e-9)

# CSV byte-identity vs the pre-load snapshot
before = dict(line.split(" ", 1)[::-1] for line in (Path(__file__).parent / "csv_hashes_before.txt").read_text().splitlines())
before = {k.strip(): v for k, v in before.items()}
changed = []
for f in sorted(IND.glob("*/*.csv")):
    rel = f.relative_to(ROOT).as_posix()
    h = subprocess.run(["git", "hash-object", rel], cwd=ROOT, capture_output=True, text=True).stdout.strip()
    if before.get(rel) != h:
        changed.append(rel)
print(f"CSV hashes {'PASS' if not changed and len(before) else 'FAIL'}: {len(before)} snapshotted, changed/new = {changed}")
if changed:
    fails.append("csv")
raise SystemExit(1 if fails else 0)
