"""Load the non-US regional rows into the eight cement input workbooks' Regions tabs.

Source: Cement_Regional_Data_Draft.md (worktree root), passes 1 and 2. Pass 2 wins on
conflict; the conflict is kept in the row note. Writes Regions rows 6-14 and one About
note only. Never touches the US row (row 5), Derivation formulas, or export tabs.

Run from anywhere:  <python> Cement_Regional_Rows_Loading/load_regional_rows.py
(the worktree root is this file's parent's parent).
Computed (JUDGMENT) cells are written as Excel formulas over the source numbers so the
arithmetic stays visible; `py` holds the Python mirror of each formula for the log.
"""
import json
from copy import copy
from pathlib import Path

import openpyxl

LOAD_DATE = "2026-10-01"
ROOT = Path(__file__).resolve().parent.parent
IND = ROOT / "InputData" / "indst"

# ---------------------------------------------------------------------------
# Step 1: region x item -> (value, year, flag, note)
# ---------------------------------------------------------------------------
D = {}


def put(region, item, value, year, flag, note, py=None):
    D.setdefault(region, {})[item] = dict(value=value, year=year, flag=flag, note=note,
                                          py=value if py is None else py)


# --- Apparent consumption anchor (Mt) -> BCD -------------------------------
put("China", "consumption", 1858.9, 2024, "SNIPPET",
    "cemnet citing demand data (pass 1, carried in pass 2). 2025 projection 1,790 (ICR via Cement Europe, SNIPPET); "
    "NBS 2025 production 1,670-1,693 (SNIPPET); 2024 fell ~8-10% y/y. Long-run CAGR TO FILL (IEA 2018 roadmap: China declines to 2050, no number).")
put("India", "consumption", 435, 2024, "SNIPPET",
    "FY2024-25 (Apr 2024-Mar 2025, assigned to 2024) consumption 435 / production 427, ACEEE Jul-2026 citing IBEF (pass 2). "
    "Pass 1 had no consumption (USGS production 440 in 2024e, 470 in 2025e - calendar vs FY basis).")
put("EU-27", "consumption", 148.1, 2024, "DIRECT",
    "Cement Europe Statistics Report 2025, EU-27 scope (pass 2 read the PDF; pass 1 had it as SNIPPET). Production 160.8 (2024, DIRECT). "
    "UNRESOLVED: production - consumption = +12.7 Mt vs ~0 Eurostat net trade - ask Cement Europe. Do not mix with the 165.1 wider-membership figure.")
put("Brazil", "consumption", 62.21, 2023, "DIRECT",
    "SNIC Annual Report 2023 apparent consumption (pass 2, DIRECT). CONFLICT: pass 1 used 64.652 (2024 SNIC sales via cemnet, SNIPPET); "
    "2024 consumption 64.8 and 2025 dispatches 66.984 are SNIPPET. A 73 Mt 2024 'consumption' figure (pass 1) is unreliable.")
put("Mexico", "consumption", 42, 2024, "SNIPPET",
    "~42 Mt 2024 sales (Cementos Moctezuma CEO via Global Cement, pass 2). CONFLICT: pass 1 had 46.4 (2024 CANACEM forecast via Global Cement) - "
    "resolve against CANACEM statistics. 2025 consumption -6% (SNIPPET).")
put("Canada", "consumption", "=12.7+1.2-(4.39+0.47)", 2024, "JUDGMENT",
    "LOW CONFIDENCE: 2020 CAC-member production 12.7 (DIRECT, Concrete Zero) + 2024 Comtrade imports 1.20 - exports (cement 4.39 + clinker 0.47) "
    "(DIRECT trade); mixes years. StatCan 16-10-0009-01 ends 2018. Next: StatCan monthly cement table / NRCan Minerals Yearbook.",
    py=round(12.7 + 1.2 - (4.39 + 0.47), 2))
put("Indonesia", "consumption", 63.912, 2025, "DIRECT",
    "Asperssi (ASI) 2025 domestic sales, used as apparent consumption (imports negligible) (pass 2, DIRECT). 2024: 64.895. "
    "Pass 1 had 64.9 (2024, SNIPPET via cemnet). Production 64.720 (2025).")

# --- Demand growth CAGR to 2050 (fraction/yr) -> BCD -----------------------
put("India", "cagr", "=ROUND((2100/391)^(1/(2070-2023))-1,3)", "2023-70", "JUDGMENT",
    "NITI Aayog Jan-2026 path 391 (2023) -> 2,100 Mt (2070) (SNIPPET); CAGR is own arithmetic. Front-loaded path: 7.8%/yr 2023-30, 5.9%/yr 2030-47; "
    "GCCA India/TERI ~3.1%/yr; IEA 2018 'triple by 2050' ~3.1%/yr. A flat CAGR understates near-term growth.",
    py=round((2100 / 391) ** (1 / (2070 - 2023)) - 1, 3))
put("Brazil", "cagr", 0.017, "2014-50", "JUDGMENT",
    "SNIC/IEA/WBCSD 2019 roadmap: demand +60% (low) to +120% (high) vs 2014 by 2050 (DIRECT) = 1.3-2.2%/yr; 1.7%/yr chosen within the range (staff plan). "
    "Applied from the 2023 anchor, which sits below the 2014 base, so 2050 lands under the roadmap low case.")
put("Canada", "cagr", 0.01, "2020-50", "DIRECT",
    "CAC Concrete Zero Action Plan assumes cement volume +1%/yr from 2020 (DIRECT).")

# --- Blend (clinker-to-cement) ratio -> BCtCR -------------------------------
put("India", "blend_ratio", 0.73, 2025, "SNIPPET",
    "ACEEE Jul-2026 citing Indian Cement Review (SNIPPET). NITI Aayog baseline 0.675 (SNIPPET); GCCA India/TERI 0.75 in 2020 (SNIPPET) - sources disagree. "
    "Production basis, not consumption blend basis (trade ~0, so close).")
put("EU-27", "blend_ratio", 0.77, 2021, "DIRECT",
    "CEMBUREAU Net Zero Roadmap text: 77% for 'Europe' in 2021 (DIRECT); Cement Europe says the ratio fell ~2% over the last 3 years. Scope is Europe, not strictly EU-27.")
put("Brazil", "blend_ratio", 0.71, 2022, "DIRECT",
    "GCCA figure in SNIC Annual Report 2023 (DIRECT); 0.84 in 1990. A snippet clinker figure (77.5 Mt 2024) exceeds cement output - unreliable.")
put("Mexico", "blend_ratio", 0.75, 2016, "DIRECT",
    "FICEM/CANACEM roadmap, PwC-verified GNR data, 2016 (DIRECT). Stale anchor year - newer CANACEM/GNR figure needed.")
put("Canada", "blend_ratio", 0.89, 2020, "DIRECT",
    "CAC Concrete Zero 2020 baseline 89% (DIRECT; 11.4 clinker / 12.7 cement = 0.90). Production basis; Canada is a net exporter.")
put("Indonesia", "blend_ratio", "=ROUND((57.301-12.34)/63.912,2)", 2025, "JUDGMENT",
    "Domestic clinker use (ASI clinker production 57.301 - clinker exports 12.34, DIRECT) / domestic sales 63.912 (DIRECT) - domestic basis. "
    "Do not use 57.3/64.7 = 0.89 (includes exported clinker). ACEEE: 0.68 in 2026 (SNIPPET).",
    py=round((57.301 - 12.34) / 63.912, 2))

# --- Net import share of clinker demand (clinker basis) -> BFoCDMbNI --------
# Clinker basis per the workbook definition: (net cement imports x ratio + net clinker imports)
# / (consumption x ratio). Trade numbers are UN Comtrade HS 2523 (DIRECT) unless noted.
put("China", "net_import_share", "=ROUND(((0.06-5.02)+(0.28-0.34))/1858.9,3)", 2024, "JUDGMENT",
    "Comtrade 2024: cement M 0.06 / X 5.02, clinker M 0.28 / X 0.34 Mt (DIRECT) over consumption 1,858.9. Tonnage basis because the China ratio is TO FILL; "
    "the clinker-basis value differs by <0.001 for any ratio 0.5-0.9. 2025 clinker exports not reported.",
    py=round(((0.06 - 5.02) + (0.28 - 0.34)) / 1858.9, 3))
put("India", "net_import_share", "=ROUND(((0.74-0.98)*0.73+(1.55-0.02))/(435*0.73),3)", 2024, "JUDGMENT",
    "Comtrade 2024: cement M 0.74 / X 0.98, clinker M 1.55 / X 0.02 Mt (DIRECT); ratio 0.73, consumption 435. Clinker basis (+0.003 on tonnage basis, pass 2). "
    "Replaces the pass 1 placeholder 0.",
    py=round(((0.74 - 0.98) * 0.73 + (1.55 - 0.02)) / (435 * 0.73), 3))
put("EU-27", "net_import_share", "=ROUND((11.3-11.2)/148.1,3)", 2024, "JUDGMENT",
    "Eurostat via Cement Europe 2024: imports 11.3 / exports 11.2 Mt cement+clinker combined (DIRECT) - tonnage basis, no cement/clinker split. "
    "2025: imports 14.2 / exports 9.1 = +0.034 of 2024 consumption. Conflicts with the +12.7 Mt production-consumption gap (see BCD note).",
    py=round((11.3 - 11.2) / 148.1, 3))
put("Mexico", "net_import_share", "=ROUND(((0.10-1.38)*0.75+0)/(42*0.75),3)", 2024, "JUDGMENT",
    "Comtrade 2024: cement M 0.10 / X 1.38 Mt, clinker ~0 (DIRECT; zero may mean not reported); ratio 0.75, consumption 42. "
    "CONFLICT: pass 1 had +0.05 (net importer) from CANACEM 46.4 vs USGS 44 - pass 2 trade data says net exporter (~1.5 Mt to US, SNIPPET).",
    py=round(((0.10 - 1.38) * 0.75 + 0) / (42 * 0.75), 3))
put("Indonesia", "net_import_share", "=ROUND(-(12.34+1.32*0.68)/(57.301-12.34),3)", 2025, "JUDGMENT",
    "ASI 2025: clinker exports 12.34, cement exports 1.32 Mt (DIRECT; Comtrade matches) at ratio 0.68 over domestic clinker use 44.96 - clinker basis. "
    "Tonnage basis -0.21. Replaces pass 1 -0.16 (mixed cement and clinker volumes).",
    py=round(-(12.34 + 1.32 * 0.68) / (57.301 - 12.34), 3))
put("Brazil", "net_import_share", "=ROUND(((0.14-0.07)*0.71+0.92)/(64.8*0.71),3)", 2024, "JUDGMENT",
    "Comtrade 2024: clinker M 0.92, cement M 0.14 / X 0.07 Mt (DIRECT); ratio 0.71, 2024 consumption 64.8 (SNIPPET). Clinker basis; tonnage basis +0.015 (2024), "
    "+0.012 (2023). Staff plan quoted +0.013.",
    py=round(((0.14 - 0.07) * 0.71 + 0.92) / (64.8 * 0.71), 3))
put("Canada", "net_import_share", "=ROUND(((1.2-4.39)*0.89-0.47)/((12.7+1.2-(4.39+0.47))*0.89),3)", 2024, "JUDGMENT",
    "LOW CONFIDENCE. Comtrade 2024: cement X 4.39, clinker X 0.47, imports 1.20 Mt (DIRECT; imports treated as cement) over the ~9.0 Mt JUDGMENT consumption, ratio 0.89. "
    "Tonnage basis -0.40; pass 2 text said 'about -35%' (range -30 to -40%) but its own -3.65/9 = -0.41. ISED: ~6 Mt/yr to the US.",
    py=round(((1.2 - 4.39) * 0.89 - 0.47) / ((12.7 + 1.2 - (4.39 + 0.47)) * 0.89), 3))

# --- Minimum achievable ratio -> MACtCR ------------------------------------
put("EU-27", "min_ratio", 0.60, 2050, "DIRECT",
    "CEMBUREAU 'From Ambition to Deployment' net-zero roadmap: 60% by 2050, revised down from 65% (DIRECT).")
put("Brazil", "min_ratio", 0.52, 2050, "DIRECT",
    "SNIC/IEA/WBCSD 2019 roadmap: 52% in 2050, 59% in 2030 (DIRECT). 2025 Net Zero Roadmap (COP30): 51% (SNIPPET).")
put("Canada", "min_ratio", 0.60, 2050, "DIRECT",
    "CAC Concrete Zero: 60% in 2050 (75% 2030, 68% 2040) (DIRECT). A snippet quotes 890 -> 630 kg/t - check the ISED roadmap figure.")
put("Mexico", "min_ratio", 0.66, 2030, "DIRECT",
    "FICEM roadmap: 66% by 2030 from 75% in 2016 (DIRECT). A 2030 figure, NOT a 2050 floor - the 2050 value is TO FILL (FICEM scenario annexes).")
put("India", "min_ratio", 0.62, 2070, "SNIPPET",
    "NITI Aayog Jan-2026 roadmap ~67.5% -> ~62% by 2070 (SNIPPET). Range 0.56-0.675: GCCA India/TERI 0.56 by 2070 (SNIPPET); ACEEE ~65% by 2030 (SNIPPET).")

# CULB/CUUB note additions (identical in both workbooks, appended to column G)
CU_NOTES = {
    "China": "PASS 2 (2026-10-01): CCA capacity 1,810 Mt + 53% utilization 2024 (SNIPPET, Global Cement) supports the 0.50 trough; still unclear if clinker or cement basis. "
             "GEM 2026 operating clinker capacity 1,636.9 Mt (DIRECT, GEM article); clinker output not found, so utilization on the GEM base is unknown.",
    "India": "PASS 2 (2026-10-01): GEM 2026 clinker capacity 402.5 Mt (DIRECT, GEM article); implied clinker utilization ~0.77 (= 0.73 x 427 / 402.5, JUDGMENT) - inside the band, "
             "but the band is cement basis (CMA).",
    "Brazil": "PASS 2 (2026-10-01): SNIC cement capacity 94 Mt/yr (DIRECT); chart prints 0.76 for 2023 vs 66.5/94 = 0.71 (JUDGMENT) - check chart definition; "
              "clinker basis ~47/61 = 0.77 (JUDGMENT, USGS capacity).",
    "Mexico": "PASS 2 (2026-10-01): one-point check only - implied clinker ~33 Mt (0.75 x 44, 2016 ratio, JUDGMENT) / USGS clinker capacity 42 = ~0.79; utilization series still TO FILL.",
    "Indonesia": "PASS 2 (2026-10-01): ASI projects 0.65 for 2025 (SNIPPET, Global Cement); clinker basis 57.3 / 83 (USGS capacity) = 0.69 (JUDGMENT).",
}

CRBW_NOTE = ("k = 1.0 is a template, not a value: k must be solved at build time from the region's start-year utilization and "
             "calibrated overhaul share (design memo Cement_Utilization_Tightness_Design.md section 4), as was done for the US.")
CTOA_NOTE = "regional import-parity premium not researched; blank = adder off (Derivation returns zero)"

ABOUT_LABEL = f"Regional rows loaded {LOAD_DATE}"
ABOUT_TEXT = ("Non-US rows of the Regions tab loaded from Cement_Regional_Data_Draft.md (cement-calculation-flow-89da5b worktree root), passes 1 and 2; "
              "pass 2 (primary sources) wins on conflict and the conflict is kept in the row note. Flags: DIRECT = read in the primary document; "
              "SNIPPET = secondary/trade-press summary or an unopened primary; JUDGMENT = own arithmetic or estimate (written as a visible formula where computed); "
              "blank = TO FILL. Supersedes any placeholder description in the Sources block above. The US row, Derivation and export tab are unchanged. "
              "Verify every non-US value against its cited primary source before use. Loader: Cement_Regional_Rows_Loading/load_regional_rows.py.")

# ---------------------------------------------------------------------------
# Steps 2-4: map items to Regions columns, write rows 6-14, add the About note.
# ---------------------------------------------------------------------------
WB = {
    "BCD": "BCD/BAU Cement Demand.xlsx",
    "BCtCR": "BCtCR/BAU Clinker to Cement Ratio.xlsx",
    "BFoCDMbNI": "BFoCDMbNI/BAU Fraction of Clinker Demand Met by Net Imports.xlsx",
    "MACtCR": "MACtCR/Minimum Achievable Clinker to Cement Ratio.xlsx",
    "CULB": "CULB/Cement Utilization Lower Bound.xlsx",
    "CUUB": "CUUB/Cement Utilization Upper Bound.xlsx",
    "CRBW": "CRBW/Cement Replacement Blend Weight vs Tightness.xlsx",
    "CTOA": "CTOA/Cement Tightness Operating Adder vs Tightness.xlsx",
}


def region_rows(ws):
    rows = {ws.cell(r, 1).value: r for r in range(5, 15)}
    assert rows.get("United States") == 5, "US row moved"
    return rows


def load_bcd(ws, rows):
    for reg, items in D.items():
        r = rows[reg]
        c, g = items.get("consumption"), items.get("cagr")
        if c:
            ws.cell(r, 2).value = c["year"]
            ws.cell(r, 3).value = c["value"]
            ws.cell(r, 4).value = c["flag"]
        if g:
            ws.cell(r, 5).value = g["value"]
            ws.cell(r, 6).value = g["flag"]
        if c or g:
            note = c["note"] if c else ws.cell(r, 8).value
            if g:
                note = f"{note} | CAGR ({g['year']}): {g['note']}"
            ws.cell(r, 8).value = note


def load_bctcr(ws, rows):
    for reg, items in D.items():
        b = items.get("blend_ratio")
        if b:
            r = rows[reg]
            ws.cell(r, 2).value = b["year"]
            ws.cell(r, 3).value = b["value"]
            ws.cell(r, 4).value = b["flag"]
            ws.cell(r, 8).value = b["note"] + " | BAU target blank = flat at anchor (no BAU trajectory sourced; roadmap targets are policy, see MACtCR)."
    ws.cell(rows["China"], 8).value = (
        "TO FILL - conflicting snippets (0.657 vs 0.735, CCA disputes the rise); pass 2 implied ~0.52 from 53% x 1.81 Bt capacity also conflicts. "
        "Next: NBS clinker output + CCA annual report.")


def load_bfocdmbni(ws, rows):
    for reg, items in D.items():
        s = items.get("net_import_share")
        if s:
            r = rows[reg]
            ws.cell(r, 2).value = s["year"]
            ws.cell(r, 3).value = s["value"]
            ws.cell(r, 4).value = s["flag"]
            ws.cell(r, 5).value = s["note"]


def load_mactcr(ws, rows):
    for reg, items in D.items():
        m = items.get("min_ratio")
        if m:
            r = rows[reg]
            ws.cell(r, 2).value = m["value"]
            ws.cell(r, 3).value = m["flag"]
            ws.cell(r, 4).value = f"({m['year']}) {m['note']}"


def load_cu(ws, rows):
    for reg, extra in CU_NOTES.items():
        r = rows[reg]
        old = ws.cell(r, 7).value or ""
        if "PASS 2" not in old:
            ws.cell(r, 7).value = f"{old} | {extra}" if old else extra


def load_crbw(ws, rows):
    for reg, r in rows.items():
        if reg != "United States":
            ws.cell(r, 5).value = CRBW_NOTE


def load_ctoa(ws, rows):
    for reg, r in rows.items():
        if reg != "United States":
            ws.cell(r, 5).value = CTOA_NOTE


LOADERS = {"BCD": load_bcd, "BCtCR": load_bctcr, "BFoCDMbNI": load_bfocdmbni, "MACtCR": load_mactcr,
           "CULB": load_cu, "CUUB": load_cu, "CRBW": load_crbw, "CTOA": load_ctoa}


def add_about_note(ws):
    """Append (or refresh) the load note as the last row of the About Notes section."""
    last = ws.max_row
    for r in range(1, last + 1):
        if ws.cell(r, 1).value == ABOUT_LABEL:
            ws.cell(r, 2).value = ABOUT_TEXT
            return
    new = last + 1
    for col in (1, 2):
        src, dst = ws.cell(last, col), ws.cell(new, col)
        dst.font, dst.alignment, dst.border, dst.fill = copy(src.font), copy(src.alignment), copy(src.border), copy(src.fill)
        dst.number_format = src.number_format
    ws.cell(new, 1).value = ABOUT_LABEL
    ws.cell(new, 2).value = ABOUT_TEXT


def main():
    for acr, rel in WB.items():
        path = IND / rel
        wb = openpyxl.load_workbook(path)
        ws = wb["Regions"]
        rows = region_rows(ws)
        us_before = [c.value for c in ws[5]]
        LOADERS[acr](ws, rows)
        assert [c.value for c in ws[5]] == us_before, f"{acr}: US row changed"
        add_about_note(wb["About"])
        wb.save(path)
        print(f"loaded {acr}: {path.name}")
    # log of loaded values (Python mirrors of formulas) for the draft's Loaded section
    log = {reg: {k: {kk: v[kk] for kk in ("py", "year", "flag")} for k, v in items.items()} for reg, items in D.items()}
    (Path(__file__).parent / "loaded_values.json").write_text(json.dumps(log, indent=1))
    print(json.dumps(log, indent=1))


if __name__ == "__main__":
    main()
